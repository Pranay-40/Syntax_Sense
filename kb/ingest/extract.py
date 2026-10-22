import io
import pypdf
import docx

class ExtractionError(Exception):
    """Raised when document extraction fails."""
    pass

def extract_txt(file_bytes):
    """
    Extract text from a plain text file bytes.
    Tries utf-8, then falls back to latin-1.
    Returns: (extracted_text, topic)
    """
    try:
        text = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        text = file_bytes.decode('latin-1', errors='replace')

    # Topic is first non-empty line
    topic = "Plain Text Document"
    for line in text.splitlines():
        clean = line.strip().lstrip('#').strip()
        if clean:
            topic = clean[:100]
            break

    return text, topic

def extract_pdf(file_bytes):
    """
    Extract text from a PDF file in-memory.
    Returns: (full_text, topic, page_count, pages_dict)
    where pages_dict is {page_number (1-indexed): text}
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        raise ExtractionError(f"Failed to read PDF file: {str(e)}")

    page_count = len(reader.pages)
    if page_count == 0:
        raise ExtractionError("PDF file contains no pages.")

    pages_dict = {}
    full_text_parts = []
    topic = None

    for idx, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        pages_dict[idx] = page_text
        if page_text.strip():
            full_text_parts.append(f"--- Page {idx} ---\n{page_text.strip()}")
            if topic is None:
                for line in page_text.splitlines():
                    clean = line.strip().lstrip('#').strip()
                    if clean:
                        topic = clean[:100]
                        break

    full_text = "\n\n".join(full_text_parts)
    if not full_text.strip():
        raise ExtractionError("PDF contains no readable text (scanned images/OCR are not supported).")

    return full_text, topic or "PDF Document", page_count, pages_dict

def extract_docx(file_bytes):
    """
    Extract text from a DOCX file in-memory.
    Returns: (full_text, topic)
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ExtractionError(f"Failed to read DOCX file: {str(e)}")

    text_parts = []
    topic = None

    for p in doc.paragraphs:
        p_text = p.text.strip()
        if p_text:
            text_parts.append(p_text)
            if topic is None:
                # First non-empty paragraph or title
                topic = p_text[:100]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    full_text = "\n\n".join(text_parts)
    if not full_text.strip():
        raise ExtractionError("DOCX file contains no text.")

    return full_text, topic or "DOCX Document"
